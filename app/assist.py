"""AI assist: turn meeting notes or an inventory spreadsheet into proposed sources and products.

* Spreadsheets (CSV / XLSX) that follow the template (GET /api/assist/template.csv)
  are mapped directly - no AI, no cost. Other layouts and free-text notes go to
  Gemini on Vertex AI in the deployment's own project, with a JSON schema that only
  allows the tool's source types, frequencies and products.
* The result is a *proposal*: the UI shows it and the architect ticks what to add.
* Model: GEMINI_MODEL (default gemini-3.8-flash), location GEMINI_LOCATION
  (default global); falls back to gemini-2.5-flash if the first model is not
  available in the project. Per-user rate limit: ASSIST_LIMIT_PER_HOUR (default 40).
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import time
from collections import defaultdict

from fastapi import HTTPException

from .catalog import FREQUENCIES, SOURCE_TYPES
from .components import COMPONENTS

TEMPLATE_COLUMNS = ["name", "type", "frequency", "daily_change_gb", "initial_volume_gb", "location", "data_type", "transform", "notes"]
LOCATIONS = ["on_prem", "gcp", "aws", "azure", "saas"]
DATA_TYPES = ["structured", "semi_structured", "unstructured", "events"]
TRANSFORMS = ["none", "light", "heavy"]
UNITS_GB = {"kb": 1e-6, "mb": 1e-3, "gb": 1.0, "tb": 1e3, "pb": 1e6, "eb": 1e9, "zb": 1e12}

_calls = defaultdict(list)


def template_csv() -> str:
    rows = [TEMPLATE_COLUMNS,
            ["Core banking", "oracle", "near_realtime", "60", "8 TB", "on_prem", "structured", "light", "Oracle 19c, ~400 tables"],
            ["Card transactions", "kafka", "realtime", "40", "0", "on_prem", "events", "light", "Confluent topics"],
            ["CRM", "salesforce", "daily", "2", "300 GB", "saas", "structured", "none", ""]]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    buf.write("\n# type: " + ", ".join(s["id"] for s in SOURCE_TYPES) + "\n")
    buf.write("# frequency: " + ", ".join(f["id"] for f in FREQUENCIES) + "\n")
    buf.write("# location: " + ", ".join(LOCATIONS) + " | data_type: " + ", ".join(DATA_TYPES) + " | transform: " + ", ".join(TRANSFORMS) + "\n")
    buf.write("# volumes: a number in GB, or with a unit (KB, MB, GB, TB, PB)\n")
    return buf.getvalue()


# --------------------------------------------------------------------------- spreadsheet parsing
def _gb(v):
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"\s*([\d.,]+)\s*([kmgtpez]i?b)?", str(v), re.I)
    if not m:
        return 0.0
    n = float(m.group(1).replace(",", ""))
    unit = (m.group(2) or "gb").lower().replace("ib", "b")
    return n * UNITS_GB.get(unit, 1.0)


def _match_type(v):
    s = str(v or "").strip().lower()
    for t in SOURCE_TYPES:
        if s == t["id"] or s == t["name"].lower():
            return t["id"]
    for t in SOURCE_TYPES:
        words = re.split(r"[^a-z0-9]+", t["name"].lower())
        if s and (s in t["name"].lower() or any(w and len(w) > 3 and w in s for w in words)):
            return t["id"]
    return None


def _enum(v, allowed, default):
    s = re.sub(r"[^a-z_]", "_", str(v or "").strip().lower()).strip("_")
    aliases = {"real_time": "realtime", "near_real_time": "near_realtime", "cdc": "near_realtime", "streaming": "realtime",
               "batch": "daily", "once": "one_time", "onprem": "on_prem", "on_premise": "on_prem", "semi": "semi_structured",
               "json": "semi_structured"}
    s = aliases.get(s, s)
    return s if s in allowed else default


def read_table(filename: str, data: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [[c for c in row] for row in ws.iter_rows(values_only=True)][:2000]
    if name.endswith((".csv", ".txt", ".tsv")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        return [r for r in csv.reader(io.StringIO(text), dialect)][:2000]
    raise HTTPException(status_code=415, detail="Upload a .csv or .xlsx file")


def from_template(rows: list[list]):
    """Rows in the template layout -> proposal, or None when the headers do not match."""
    rows = [r for r in rows if r and any(c not in (None, "") for c in r) and not str(r[0] or "").startswith("#")]
    if not rows:
        return None
    head = [str(h or "").strip().lower().replace(" ", "_") for h in rows[0]]
    if "name" not in head or "type" not in head:
        return None
    idx = {k: head.index(k) for k in TEMPLATE_COLUMNS if k in head}
    freqs = [f["id"] for f in FREQUENCIES]
    sources, problems = [], []
    for n, r in enumerate(rows[1:], start=2):
        get = lambda k: r[idx[k]] if k in idx and idx[k] < len(r) else None  # noqa: E731
        t = _match_type(get("type"))
        if not t:
            problems.append(f"Row {n}: unknown type '{get('type')}' - skipped")
            continue
        sources.append({"name": str(get("name") or t)[:120], "type": t,
                        "frequency": _enum(get("frequency"), freqs, "daily"),
                        "daily_change_gb": _gb(get("daily_change_gb")), "initial_volume_gb": _gb(get("initial_volume_gb")),
                        "location": _enum(get("location"), LOCATIONS, "on_prem"),
                        "data_type": _enum(get("data_type"), DATA_TYPES, "structured"),
                        "transform": _enum(get("transform"), TRANSFORMS, "none"),
                        "notes": str(get("notes") or "")[:2000]})
    return {"sources": sources, "components": [], "questions": [], "assumptions": problems, "method": "template"}


# --------------------------------------------------------------------------- Gemini
def _schema():
    costed = [c["id"] for c in COMPONENTS]
    return {
        "type": "OBJECT",
        "properties": {
            "customer": {"type": "STRING"},
            "industry": {"type": "STRING"},
            "region_hint": {"type": "STRING", "description": "Google Cloud region id if data residency is stated, e.g. me-central2"},
            "sources": {"type": "ARRAY", "items": {"type": "OBJECT", "properties": {
                "name": {"type": "STRING"},
                "type": {"type": "STRING", "enum": [s["id"] for s in SOURCE_TYPES]},
                "frequency": {"type": "STRING", "enum": [f["id"] for f in FREQUENCIES]},
                "daily_change_gb": {"type": "NUMBER"},
                "initial_volume_gb": {"type": "NUMBER"},
                "location": {"type": "STRING", "enum": LOCATIONS},
                "data_type": {"type": "STRING", "enum": DATA_TYPES},
                "transform": {"type": "STRING", "enum": TRANSFORMS},
                "notes": {"type": "STRING", "description": "What the notes said about this system; mark guessed volumes as 'estimated'"},
                "volume_estimated": {"type": "BOOLEAN"},
            }, "required": ["name", "type", "frequency", "daily_change_gb", "initial_volume_gb", "location"]}},
            "components": {"type": "ARRAY", "items": {"type": "OBJECT", "properties": {
                "type": {"type": "STRING", "enum": costed},
                "reason": {"type": "STRING"},
                "users": {"type": "NUMBER"},
            }, "required": ["type", "reason"]}},
            "environments": {"type": "ARRAY", "items": {"type": "STRING", "enum": ["prod", "preprod", "uat", "dev", "sandbox", "dr"]}},
            "questions": {"type": "ARRAY", "items": {"type": "STRING"}, "description": "Open questions to ask the customer next"},
            "assumptions": {"type": "ARRAY", "items": {"type": "STRING"}},
        },
        "required": ["sources", "components", "questions", "assumptions"],
    }


PROMPT = """You are a Google Cloud data & AI architect's assistant in a customer discovery session.
From the material below, list the source systems that must land in BigQuery and the analytics, AI, API,
database, governance and security products the customer mentioned or clearly needs.
Rules:
- Use only the allowed enum values. Map each system to the closest source type.
- Volumes in GB: daily_change_gb is the data changed or added per day; initial_volume_gb the historical load.
  If not stated, estimate conservatively from typical sizes, set volume_estimated=true and say so in notes.
- Frequency: realtime (seconds), near_realtime (minutes / CDC), hourly, daily, weekly, monthly, one_time (migration only).
- Location: where the system runs today.
- Only propose products that the material mentions or clearly implies; give a one-line reason each.
- Add the 3-8 most useful open questions to ask the customer next, and list assumptions you made.
Material:
"""


def _check_rate(email):
    limit = int(os.environ.get("ASSIST_LIMIT_PER_HOUR", "40"))
    now = time.time()
    calls = [t for t in _calls[email] if now - t < 3600]
    if len(calls) >= limit:
        raise HTTPException(status_code=429, detail=f"AI assist limit of {limit} requests per hour reached - try again later")
    calls.append(now)
    _calls[email] = calls


def gemini(text: str, email: str) -> dict:
    if os.environ.get("ASSIST", "on") == "off":
        raise HTTPException(status_code=503, detail="AI assist is turned off on this deployment")
    _check_rate(email)
    try:
        from google import genai
        from google.genai import types
    except ImportError as e:
        raise HTTPException(status_code=503, detail="AI assist is not installed on this deployment") from e
    project = os.environ.get("GOOGLE_CLOUD_PROJECT") or os.environ.get("GCP_PROJECT")
    location = os.environ.get("GEMINI_LOCATION", "global")
    models = [m for m in (os.environ.get("GEMINI_MODEL", "gemini-3.8-flash"), "gemini-2.5-flash") if m]
    client = genai.Client(vertexai=True, project=project, location=location)
    cfg = types.GenerateContentConfig(response_mime_type="application/json", response_schema=_schema(), temperature=0.2,
                                      max_output_tokens=8192)
    last = None
    for model in dict.fromkeys(models):
        try:
            r = client.models.generate_content(model=model, contents=PROMPT + text[:60_000], config=cfg)
            data = json.loads(r.text)
            data["method"] = "gemini"
            data["model"] = model
            return clean(data)
        except Exception as e:  # noqa: BLE001 - try the fallback model, then report
            last = e
    raise HTTPException(status_code=502, detail=f"Gemini could not process the request: {type(last).__name__}: {str(last)[:300]}")


def clean(data: dict) -> dict:
    """Keep only valid values (the schema should guarantee it - this is the safety net)."""
    types_ = {s["id"] for s in SOURCE_TYPES}
    freqs = {f["id"] for f in FREQUENCIES}
    comps = {c["id"] for c in COMPONENTS}
    out = {k: data.get(k) for k in ("customer", "industry", "region_hint", "method", "model") if data.get(k)}
    out["sources"] = []
    for s in data.get("sources") or []:
        if s.get("type") not in types_:
            continue
        out["sources"].append({
            "name": str(s.get("name") or s["type"])[:120], "type": s["type"],
            "frequency": s.get("frequency") if s.get("frequency") in freqs else "daily",
            "daily_change_gb": max(0.0, float(s.get("daily_change_gb") or 0)), "initial_volume_gb": max(0.0, float(s.get("initial_volume_gb") or 0)),
            "location": s.get("location") if s.get("location") in LOCATIONS else "on_prem",
            "data_type": s.get("data_type") if s.get("data_type") in DATA_TYPES else "structured",
            "transform": s.get("transform") if s.get("transform") in TRANSFORMS else "none",
            "notes": str(s.get("notes") or "")[:2000], "volume_estimated": bool(s.get("volume_estimated")),
        })
    seen = set()
    out["components"] = []
    for c in data.get("components") or []:
        if c.get("type") in comps and c["type"] not in seen:
            seen.add(c["type"])
            out["components"].append({"type": c["type"], "reason": str(c.get("reason") or "")[:300],
                                      **({"users": int(c["users"])} if c.get("users") else {})})
    out["environments"] = [e for e in (data.get("environments") or []) if e in ("prod", "preprod", "uat", "dev", "sandbox", "dr")]
    out["questions"] = [str(q)[:300] for q in (data.get("questions") or [])][:12]
    out["assumptions"] = [str(q)[:300] for q in (data.get("assumptions") or [])][:20]
    return out


def table_as_text(rows: list[list]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in rows[:400]:
        w.writerow(["" if c is None else str(c) for c in r])
    return "Inventory spreadsheet (CSV):\n" + buf.getvalue()
